"""List PowerPoint decks in a folder that have no summary yet, or extract one deck's text.

Usage:
    python find_new_ppts.py <training_folder>
    python find_new_ppts.py <training_folder> --extract <deck_path>
"""
import argparse
import sys
from pathlib import Path

PPT_EXTENSIONS = {".pptx", ".ppt"}
SUMMARY_DIR = "Summaries"


def find_new(folder: Path) -> list[Path]:
    summaries = folder / SUMMARY_DIR
    decks = sorted(
        p for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in PPT_EXTENSIONS and not p.name.startswith("~$")
    )
    return [p for p in decks if not (summaries / f"{p.stem}_summary.md").exists()]


def extract(deck: Path) -> str:
    if deck.suffix.lower() == ".ppt":
        return f"{deck.name}: legacy .ppt format - save it as .pptx in PowerPoint first."
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is not installed. Run: pip install python-pptx")

    prs = Presentation(str(deck))
    lines = [f"# {deck.name} ({len(prs.slides)} slides)"]
    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title and slide.shapes.title.has_text_frame else ""
        lines.append(f"\n## Slide {i}: {title or '(no title)'}")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(f"{'  ' * para.level}- {text}")
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("folder", type=Path)
    parser.add_argument("--extract", type=Path, help="deck to extract text from")
    args = parser.parse_args()

    if not args.folder.is_dir():
        sys.exit(f"Folder not found: {args.folder}")
    if args.extract:
        print(extract(args.extract))
        return
    for deck in find_new(args.folder):
        print(deck)


if __name__ == "__main__":
    main()
